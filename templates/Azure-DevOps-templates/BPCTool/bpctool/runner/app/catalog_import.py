"""Build catalog work-item drafts from Excel with HTML/markdown normalization."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .excel_ir import _cell, is_catalog_source_file
from .mapping import TemplateMapping, normalize
from .richtext import normalize_field_value, normalize_rich_text

WORK_ITEM_TYPE_ALIASES = {
    "test case": "Test Case",
    "testcase": "Test Case",
    "process area": "Process area",
    "functional area": "Functional area",
    "system process": "System process",
    "end to end": "End to end",
}


@dataclass
class WorkItemDraft:
    key: str
    parent_key: str | None
    work_item_type: str
    title: str
    fields: dict[str, Any]
    source_file: str
    source_row: int
    hierarchy_level: int | None = None
    content_formats: dict[str, str] = field(default_factory=dict)
    # Field refs that should be created with /multilineFieldsFormat = Markdown
    markdown_format_fields: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    # Set when another row already used this row's key; ``key`` was made unique.
    duplicate_of: str = ""


def load_catalog_drafts(
    catalog_dir: str | Path | None,
    mapping: TemplateMapping,
    *,
    limit: int | None = None,
    convert_markdown_to_html: bool = False,
) -> list[WorkItemDraft]:
    if not catalog_dir:
        return []
    root = Path(catalog_dir).expanduser().resolve()
    if not root.exists():
        raise FileNotFoundError(f"Catalog source folder not found: {root}")

    drafts: list[WorkItemDraft] = []
    seen_keys: set[str] = set()
    for path in sorted(root.glob("*.xlsx")):
        if not is_catalog_source_file(path):
            continue
        drafts.extend(
            _drafts_from_workbook(
                path,
                mapping,
                convert_markdown_to_html=convert_markdown_to_html,
                seen_keys=seen_keys,
            )
        )
        if limit is not None and len(drafts) >= limit:
            drafts = drafts[:limit]
            break
    # Parents come from Title 1..N outline (ADO CSV import style), not process
    # sequence IDs. See Microsoft Learn: about-import-catalog-devops.
    return sort_parent_first(drafts)


def _drafts_from_workbook(
    path: Path,
    mapping: TemplateMapping,
    *,
    convert_markdown_to_html: bool = False,
    seen_keys: set[str] | None = None,
) -> list[WorkItemDraft]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []

    header_idx = 0
    if len(rows) > 1 and any(_cell(v).lower() == "work item type" for v in rows[1]):
        header_idx = 1
    headers = [_cell(v) for v in rows[header_idx]]
    drafts: list[WorkItemDraft] = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        values = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        if not any(v not in (None, "") for v in values.values()):
            continue
        values["_source_file"] = path.name
        values["_source_sheet"] = sheet_name
        values["_source_row"] = offset
        drafts.append(
            _row_to_draft(
                values, mapping, convert_markdown_to_html=convert_markdown_to_html
            )
        )
    _make_keys_unique(drafts, set() if seen_keys is None else seen_keys)
    # Parent/child links follow Title 1..N columns (ADO CSV hierarchy), matching
    # https://learn.microsoft.com/dynamics365/guidance/business-processes/about-import-catalog-devops
    _fill_title_hierarchy_parents(drafts)
    return drafts


def _make_keys_unique(drafts: list[WorkItemDraft], seen: set[str]) -> None:
    """Give rows that repeat an earlier row's key (e.g. Process sequence ID) their own key.

    The first row keeps the plain key so existing ``ado-id-map.csv`` entries still match.
    Later rows use ``<key>#<Microsoft ID>`` (stable across releases), else the source row.
    """
    for draft in drafts:
        if draft.key not in seen:
            seen.add(draft.key)
            continue
        original = draft.key
        microsoft_id = str(draft.fields.get("MSBPC.microsoftid") or "").strip()
        suffix = microsoft_id if microsoft_id and microsoft_id != original else ""
        candidate = f"{original}#{suffix}" if suffix else f"{original}#{draft.source_file}:{draft.source_row}"
        while candidate in seen:
            candidate = f"{candidate}+"
        draft.key = candidate
        draft.duplicate_of = original
        seen.add(candidate)


def _row_to_draft(
    row: dict[str, Any],
    mapping: TemplateMapping,
    *,
    convert_markdown_to_html: bool = False,
) -> WorkItemDraft:
    lookup = _build_lookup(row)
    wit = _work_item_type(lookup, row)
    hierarchy_title, hierarchy_level = _title_from_row(lookup)
    title = hierarchy_title or _first_value(
        lookup,
        "title",
        "name",
        "work item title",
        "process name",
        "business process",
        "deliverable",
        "test case",
        "description",
    )
    key = _external_key(lookup, row)
    if not title:
        title = f"{wit} {key}"

    fields: dict[str, Any] = {"System.Title": title}
    formats: dict[str, str] = {}
    markdown_format_fields: list[str] = []
    description = _first_value(lookup, "description")
    if description:
        is_html = _looks_like_html_marker(description)
        fields["System.Description"] = normalize_rich_text(
            description,
            field_type="html",
            convert_markdown_to_html=convert_markdown_to_html,
        )
        if is_html:
            formats["System.Description"] = "html"
        elif convert_markdown_to_html:
            formats["System.Description"] = "markdown→html"
        else:
            formats["System.Description"] = "markdown"
            markdown_format_fields.append("System.Description")
    # No Description in the spreadsheet → omit System.Description (do not invent one).

    for header, value in row.items():
        if str(header).startswith("_") or value in (None, ""):
            continue
        ref = mapping.field_ref_for_header(str(header))
        if not ref or ref in {"System.Title", "System.Description"}:
            continue
        fd = mapping.field_def(ref)
        field_type = fd.field_type if fd else ("html" if mapping.is_html_field(ref) else "")
        if mapping.is_html_field(ref) or "html" in field_type.lower():
            is_html = _looks_like_html_marker(value)
            fields[ref] = normalize_field_value(
                value,
                field_type or "html",
                convert_markdown_to_html=convert_markdown_to_html,
            )
            if is_html:
                formats[ref] = "html"
            elif convert_markdown_to_html:
                formats[ref] = "markdown→html"
            else:
                formats[ref] = "markdown"
                markdown_format_fields.append(ref)
        else:
            fields[ref] = normalize_field_value(
                value, field_type, convert_markdown_to_html=convert_markdown_to_html
            )

    fields.setdefault("MSBPC.microsoftid", _first_value(lookup, "microsoft id") or key)
    fields.setdefault("MSBPC.processsequenceid", _first_value(lookup, "process sequence id") or key)

    return WorkItemDraft(
        key=key,
        parent_key=_parent_key(lookup),
        work_item_type=wit,
        title=title,
        fields={ref: value for ref, value in fields.items() if value not in (None, "")},
        source_file=str(row.get("_source_file") or ""),
        source_row=int(row.get("_source_row") or 0),
        hierarchy_level=hierarchy_level,
        content_formats=formats,
        markdown_format_fields=markdown_format_fields,
    )


def _looks_like_html_marker(value: Any) -> bool:
    from .richtext import looks_like_html

    return looks_like_html(value)


def sort_parent_first(drafts: list[WorkItemDraft]) -> list[WorkItemDraft]:
    by_key = {d.key: d for d in drafts}
    ordered: list[WorkItemDraft] = []
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(draft: WorkItemDraft) -> None:
        if draft.key in visited:
            return
        if draft.key in visiting:
            draft.warnings.append("Cycle detected in parent relationship.")
            return
        visiting.add(draft.key)
        if draft.parent_key and draft.parent_key in by_key:
            visit(by_key[draft.parent_key])
        visiting.remove(draft.key)
        visited.add(draft.key)
        ordered.append(draft)

    for draft in drafts:
        visit(draft)
    return ordered


def _build_lookup(row: dict[str, Any]) -> dict[str, Any]:
    lookup: dict[str, Any] = {}
    for header, value in row.items():
        key = normalize(header)
        lookup[key] = value
    return lookup


def _title_from_row(lookup: dict[str, Any]) -> tuple[str | None, int | None]:
    titles: list[tuple[int, str]] = []
    for index in range(1, 11):
        value = _first_value(lookup, f"title {index}")
        if value:
            titles.append((index, value))
    if titles:
        return titles[-1][1], titles[-1][0]
    return None, None


def _work_item_type(lookup: dict[str, Any], row: dict[str, Any]) -> str:
    raw = _first_value(lookup, "work item type", "type", "ado work item type")
    if raw:
        return WORK_ITEM_TYPE_ALIASES.get(normalize(raw), str(raw).strip())
    file_name = normalize(row.get("_source_file"))
    if "deliver" in file_name:
        return "Deliverable"
    if "objective" in file_name:
        return "Objective"
    return "System process"


def _external_key(lookup: dict[str, Any], row: dict[str, Any]) -> str:
    value = _first_value(
        lookup,
        "process sequence id",
        "microsoft id",
        "partner id",
        "mavim id",
        "id",
        "external id",
        "key",
    )
    if value:
        return str(value).strip()
    stem = Path(str(row.get("_source_file") or "source")).stem
    return f"{stem}:{row.get('_source_sheet') or ''}:{row.get('_source_row')}"


def _parent_key(lookup: dict[str, Any]) -> str | None:
    """Only honor explicit parent columns; hierarchy otherwise comes from Titles."""
    explicit = _first_value(
        lookup,
        "parent key",
        "parent id",
        "parent microsoft id",
        "parent process sequence id",
        "parent",
    )
    if explicit:
        return str(explicit).strip()
    return None


def _fill_title_hierarchy_parents(drafts: list[WorkItemDraft]) -> None:
    """Set parent_key from Title 1..N columns (ADO CSV import hierarchy).

    Same model as Azure DevOps bulk CSV import and the Business Process Catalog
    guidance: Title 1 is the root, Title 2 a child of the prior Title 1, etc.
    Process sequence ID is stored as a field/key but is not used for linking.
    Explicit parent_* columns on a row win over the outline.
    """
    last_at_level: dict[int, str] = {}
    for draft in drafts:
        level = draft.hierarchy_level
        if level is None:
            continue
        if not draft.parent_key:
            for candidate in range(level - 1, 0, -1):
                if candidate in last_at_level:
                    draft.parent_key = last_at_level[candidate]
                    break
        for existing in list(last_at_level):
            if existing >= level:
                del last_at_level[existing]
        last_at_level[level] = draft.key


def _first_value(lookup: dict[str, Any], *keys: str) -> str | None:
    for key in keys:
        value = lookup.get(normalize(key))
        if value not in (None, ""):
            return str(value).strip()
    return None
