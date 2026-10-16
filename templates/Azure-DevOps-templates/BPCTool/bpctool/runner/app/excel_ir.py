"""Parse the quarterly ADO template guideline workbook (+ optional catalog sources) into IR."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

IR_VERSION = "2026.10-1"

# Release suffixes vary each quarter, e.g. "ADO template guideline (<Month>).xlsx".
TEMPLATE_FILE_GLOB = "ADO template guideline*.xlsx"

REQUIRED_TEMPLATE_SHEETS = [
    "Work item types",
    "Fields",
    "Picklists",
    "Area paths",
    "Teams",
    "Backlogs",
    "Iteration paths",
]

CATALOG_FILE_HINTS = (
    "business process catalog",
    "deliverables",
    "objectives",
    "delivery plan",
    "sbd",
)

# Template / guideline workbooks must never be treated as catalog sources even if
# their names carry the same quarterly release suffix as the catalog files.
CATALOG_EXCLUDE_HINTS = (
    "ado template",
    "template guideline",
    "guideline",
)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {path.name}")
        ws = wb[sheet_name]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _header_map(header_row: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, value in enumerate(header_row):
        key = _cell(value)
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def _get(row: list[Any], headers: dict[str, int], *names: str) -> str:
    for name in names:
        index = headers.get(name)
        if index is not None and index < len(row):
            return _cell(row[index])
    return ""


def workbook_fingerprint(path: Path, version_label: str = "") -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    label = (version_label or "unversioned").strip().lower()
    return f"sha256:{digest.hexdigest()[:16]}::{label}"


_WIT_META_HEADERS = {
    "work item type",
    "custom work item type",
    "inherit from",
    "reference name",
    "color",
    "icon",
    "help text",
    "purpose",
    "backlog type",
    "backlog name",
}


def parse_work_item_types(path: Path) -> list[dict[str, Any]]:
    rows = _sheet_rows(path, "Work item types")
    if not rows:
        return []
    header_values = [_cell(value) for value in rows[0]]
    headers = _header_map(rows[0])
    field_cols = [
        (index, name)
        for index, name in enumerate(header_values)
        if name and name.strip().lower() not in _WIT_META_HEADERS
    ]
    items: list[dict[str, Any]] = []
    for row in rows[1:]:
        name = _get(row, headers, "Work item type")
        if not name:
            continue
        custom = _get(row, headers, "Custom work item type").lower()
        # X = field on the WIT and its form; H = field on the WIT, hidden from the form.
        applicable = []
        hidden = []
        for index, label in field_cols:
            flag = _cell(row[index]).upper() if index < len(row) else ""
            if flag == "X":
                applicable.append(label)
            elif flag == "H":
                hidden.append(label)
        items.append(
            {
                "name": name,
                "custom": custom in {"yes", "true", "1"},
                "disabled": custom == "disabled",
                "inherit_from": _get(row, headers, "Inherit from"),
                "reference_name": _get(row, headers, "Reference name"),
                "backlog_type": _get(row, headers, "Backlog type"),
                "backlog_name": _get(row, headers, "Backlog name"),
                "color": _get(row, headers, "Color") or "0078D4",
                "icon": _get(row, headers, "Icon") or "icon_gear",
                "help_text": _get(row, headers, "Help text"),
                "applicable_field_labels": applicable,
                "hidden_field_labels": hidden,
            }
        )
    return items


def parse_fields(path: Path) -> list[dict[str, Any]]:
    rows = _sheet_rows(path, "Fields")
    if not rows:
        return []
    headers = _header_map(rows[0])
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows[1:]:
        reference = _get(row, headers, "Reference name")
        field_name = _get(row, headers, "Field name", "Label")
        key = reference or field_name
        if not key or key in seen:
            continue
        seen.add(key)
        field_type = _get(row, headers, "Field type")
        multi = _get(row, headers, "Multi-select field")
        items.append(
            {
                "field_name": field_name,
                "label": _get(row, headers, "Label", "Field name"),
                "reference_name": reference,
                "field_type": field_type,
                "page_name": _get(row, headers, "Page name"),
                "group_name": _get(row, headers, "Group name"),
                "group_location": _get(row, headers, "Group location"),
                "group_sequence": _get(row, headers, "Group sequence"),
                "field_sequence": _get(row, headers, "Field sequence"),
                "description": _get(row, headers, "Description"),
                "default_value": _get(row, headers, "Default value"),
                "custom_field": _get(row, headers, "Custom field").lower() in {"yes", "true", "1"},
                "required": _get(row, headers, "Required"),
                "multi_select": multi.lower() in {"yes", "true", "1", "x"},
                "is_html": "html" in field_type.lower(),
            }
        )
    return items


def parse_picklists(path: Path) -> dict[str, list[str]]:
    rows = _sheet_rows(path, "Picklists")
    if not rows:
        return {}
    headers = [_cell(value) for value in rows[0]]
    lists: dict[str, list[str]] = {name: [] for name in headers if name}
    for row in rows[1:]:
        for index, name in enumerate(headers):
            if not name or index >= len(row):
                continue
            value = _cell(row[index])
            if value and value not in lists[name]:
                lists[name].append(value)
    return {name: values for name, values in lists.items() if values}


def parse_iteration_paths(path: Path) -> list[dict[str, str]]:
    rows = _sheet_rows(path, "Iteration paths")
    if not rows:
        return []
    headers = _header_map(rows[0])
    items: list[dict[str, str]] = []
    current_l1 = ""
    for row in rows[1:]:
        level1 = _get(row, headers, "Level 1")
        level2 = _get(row, headers, "Level 2")
        if level1:
            current_l1 = level1
            items.append({"path": level1, "level": "1"})
        if level2 and current_l1:
            items.append({"path": f"{current_l1}\\{level2}", "level": "2"})
        elif level2:
            items.append({"path": level2, "level": "2"})
    return items


def parse_area_paths(path: Path) -> list[dict[str, Any]]:
    """Parse Area paths sheet using L1–L4 semantics.

    Column A is the project-name placeholder (skipped). Columns B–E are
    L1–L4 hierarchy; column F is Teams.
    """
    rows = _sheet_rows(path, "Area paths")
    if len(rows) < 2:
        return []

    team_col = None
    for index, value in enumerate(rows[0]):
        if _cell(value).lower() == "teams":
            team_col = index
            break
    # The template uses B..E as L1..L4 (indexes 1..4) when the sheet has the standard layout.
    if team_col is None:
        team_col = 5
    level_cols = list(range(1, team_col))[:4]
    if not level_cols:
        # Fallback: treat all columns before Teams as hierarchy (legacy).
        level_cols = list(range(0, team_col))

    items: list[dict[str, Any]] = []
    current: list[str] = []
    for row in rows[1:]:
        team = _cell(row[team_col]) if team_col < len(row) else ""
        # Find the deepest level cell provided on this row.
        filled_at = None
        for offset, col in enumerate(level_cols):
            value = _cell(row[col]) if col < len(row) else ""
            if value:
                # Truncate deeper levels when a shallower cell is set.
                current = current[:offset]
                current.append(value)
                filled_at = offset
        if filled_at is None or not current:
            continue
        items.append(
            {
                "path": "\\".join(current),
                "depth": filled_at,
                "team": team,
                "default_segment": current[filled_at],
            }
        )
    return items


PROJECT_NAME_PLACEHOLDER_RE = re.compile(
    r"[<\[{(]?\s*your\s+project\s+name\s+here\s*[>\]})]?", re.IGNORECASE
)


def replace_project_placeholder(value: Any, project_name: str) -> Any:
    """Swap the workbook's "Your Project Name Here" placeholder for the target project."""
    if not isinstance(value, str) or not project_name.strip():
        return value
    return PROJECT_NAME_PLACEHOLDER_RE.sub(lambda _: project_name.strip(), value)


def _apply_project_name(
    areas: list[dict[str, Any]],
    iterations: list[dict[str, str]],
    teams: list[dict[str, str]],
    project_name: str,
) -> int:
    """Replace the placeholder in place; returns the number of rows changed."""
    replaced = 0
    for items, keys in (
        (areas, ("path", "default_segment")),
        (iterations, ("path",)),
        (teams, ("backlog_iteration",)),
    ):
        for item in items:
            changed = False
            for key in keys:
                new = replace_project_placeholder(item.get(key), project_name)
                if new != item.get(key):
                    item[key] = new
                    changed = True
            replaced += changed
    return replaced


def parse_teams(path: Path) -> list[dict[str, str]]:
    rows = _sheet_rows(path, "Teams")
    if not rows:
        return []
    headers = _header_map(rows[0])
    items: list[dict[str, str]] = []
    for row in rows[1:]:
        name = _get(row, headers, "Teams", "Team")
        if not name:
            continue
        items.append(
            {
                "name": name,
                "bug_behavior": _get(row, headers, "Bug behavior"),
                "include_sub_areas": _get(row, headers, "Include sub areas"),
                "backlog_iteration": _get(row, headers, "Backlog iteration"),
            }
        )
    return items


def parse_backlogs(path: Path) -> list[dict[str, str]]:
    rows = _sheet_rows(path, "Backlogs")
    if not rows:
        return []
    headers = _header_map(rows[0])
    items: list[dict[str, str]] = []
    for row in rows[1:]:
        name = _get(row, headers, "Backlog name")
        if not name:
            continue
        items.append(
            {
                "name": name,
                "type": _get(row, headers, "Backlog type"),
                "color": _get(row, headers, "Color"),
                "default_wit": _get(row, headers, "Default work item type"),
                "rename_from": _get(row, headers, "Rename from"),
            }
        )
    return items


def _looks_like_catalog(path: Path) -> bool:
    lowered = path.name.lower()
    return any(hint in lowered for hint in CATALOG_FILE_HINTS)


def is_catalog_source_file(path: Path) -> bool:
    """True when an .xlsx in the catalog folder should be imported as catalog data.

    Accepts names with any release suffix, e.g. ``Deliverables <MMM YYYY>.xlsx``
    (hint match). Rejects the
    ADO template guideline workbook even when it sits in the same folder.
    """
    name = path.name
    if not name.lower().endswith(".xlsx") or name.startswith("~$"):
        return False
    lowered = name.lower()
    if any(hint in lowered for hint in CATALOG_EXCLUDE_HINTS):
        return False
    return _looks_like_catalog(path)


_CATALOG_KINDS = (
    ("business process catalog", "Business Process Catalog"),
    ("deliverables", "Deliverables"),
    ("objectives", "Objectives"),
    ("delivery plan", "SbD Delivery Plan"),
    ("sbd", "SbD Delivery Plan"),
)


def catalog_kind(path: Path) -> str:
    lowered = path.name.lower()
    for hint, kind in _CATALOG_KINDS:
        if hint in lowered:
            return kind
    return "Other"


def catalog_folder_warnings(catalog_dir: Path, files: list[dict[str, Any]]) -> list[str]:
    """Layout problems that would import the wrong data (mixed or missing releases)."""
    warnings: list[str] = []
    by_kind: dict[str, list[str]] = {}
    for item in files:
        by_kind.setdefault(item["kind"], []).append(item["file_name"])
    for kind, names in sorted(by_kind.items()):
        if len(names) > 1:
            warnings.append(
                f"Catalog folder has {len(names)} {kind} workbooks ({', '.join(names)}). "
                "All of them will be imported — keep one quarterly release per folder "
                "to avoid duplicate work items."
            )
    if not files:
        nested = sorted(
            child.name
            for child in catalog_dir.iterdir()
            if child.is_dir()
            and not child.name.startswith((".", "_"))
            and any(is_catalog_source_file(p) for p in child.glob("*.xlsx"))
        )
        if nested:
            warnings.append(
                "Catalog folder has no catalog workbooks of its own but contains release "
                f"subfolders ({', '.join(nested)}). Point Catalog source folder at one release "
                "folder (Load runner defaults picks the newest)."
            )
    return warnings


def find_template_workbook(folder: Path) -> Path | None:
    """Most recently modified ``ADO template guideline*.xlsx`` in ``folder``."""
    if not folder.is_dir():
        return None
    candidates = [
        path
        for path in folder.glob(TEMPLATE_FILE_GLOB)
        if path.is_file() and not path.name.startswith("~$")
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda path: (path.stat().st_mtime, path.name))


def parse_catalog_sources(catalog_dir: Path | None) -> dict[str, Any]:
    if catalog_dir is None:
        return {"files": [], "total_rows": 0, "work_item_types": {}}
    if not catalog_dir.exists() or not catalog_dir.is_dir():
        raise FileNotFoundError(f"Catalog source folder not found: {catalog_dir}")

    files: list[dict[str, Any]] = []
    type_counts: dict[str, int] = {}
    total_rows = 0

    for path in sorted(catalog_dir.glob("*.xlsx")):
        if not is_catalog_source_file(path):
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_name = wb.sheetnames[0]
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()

        if not rows:
            continue

        # Prefer second header row when present (catalog workbooks may have dual header lines).
        header_idx = 0
        if len(rows) > 1 and any(_cell(v).lower() == "work item type" for v in rows[1]):
            header_idx = 1
        headers = _header_map(list(rows[header_idx]))
        wit_col = headers.get("Work item type")
        data_rows = rows[header_idx + 1 :]
        row_count = 0
        sample_titles: list[str] = []
        for row in data_rows:
            values = [_cell(v) for v in row]
            if not any(values):
                continue
            row_count += 1
            if wit_col is not None and wit_col < len(row):
                wit = _cell(row[wit_col]) or "(blank)"
                type_counts[wit] = type_counts.get(wit, 0) + 1
            if len(sample_titles) < 3:
                # Title 1 often holds the hierarchy root/label
                title = ""
                for key in ("Title 1", "Title", "Name"):
                    if key in headers and headers[key] < len(row):
                        title = _cell(row[headers[key]])
                        if title:
                            break
                if title:
                    sample_titles.append(title)

        total_rows += row_count
        files.append(
            {
                "file_name": path.name,
                "path": str(path),
                "kind": catalog_kind(path),
                "sheet": sheet_name,
                "row_count": row_count,
                "sample_titles": sample_titles,
            }
        )

    return {
        "files": files,
        "total_rows": total_rows,
        "work_item_types": dict(sorted(type_counts.items(), key=lambda item: (-item[1], item[0]))),
        "warnings": catalog_folder_warnings(catalog_dir, files),
    }


def validate_template_sheets(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        present = set(wb.sheetnames)
    finally:
        wb.close()
    return [name for name in REQUIRED_TEMPLATE_SHEETS if name not in present]


def build_ir(
    excel_file: str,
    *,
    catalog_source_dir: str | None = None,
    workbook_version: str = "",
    process_name: str = "",
    project_name: str = "",
) -> dict[str, Any]:
    raw = Path(excel_file).expanduser()
    if not excel_file or not str(excel_file).strip():
        raise FileNotFoundError("Excel template path is empty.")
    if not raw.is_absolute():
        raise FileNotFoundError(
            f'Excel template path must be absolute (got "{excel_file}"). '
            "Use Browse… in the UI (local runner path picker) or paste a full path "
            "such as /Users/you/.../ADO template guideline.xlsx — "
            "browser dialogs alone often return only the folder/file name."
        )
    path = raw.resolve()
    if path.is_dir():
        raise FileNotFoundError(
            f'Excel template must be an .xlsx file, not a folder: {path}. '
            "Point Template workbook path at the guideline workbook; use Catalog source "
            "folder for the BPC Files directory."
        )
    if not path.exists():
        raise FileNotFoundError(f"Excel template not found: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Excel template must be an .xlsx file")

    missing = validate_template_sheets(path)
    if missing:
        raise ValueError("Excel template is missing required sheets: " + ", ".join(missing))

    catalog_dir = (
        Path(catalog_source_dir).expanduser().resolve() if catalog_source_dir else None
    )
    catalog = parse_catalog_sources(catalog_dir)

    # Keep disabled WITs in IR so process can disable them in ADO and backlog
    # assignment can skip them (Work item types sheet).
    work_item_types = parse_work_item_types(path)
    fields = parse_fields(path)
    picklists = parse_picklists(path)
    iterations = parse_iteration_paths(path)
    areas = parse_area_paths(path)
    teams = parse_teams(path)
    placeholders_replaced = _apply_project_name(areas, iterations, teams, project_name or "")
    backlogs = parse_backlogs(path)
    html_fields = [item for item in fields if item.get("is_html")]
    enabled_wits = [item for item in work_item_types if not item.get("disabled")]
    disabled_wits = [item for item in work_item_types if item.get("disabled")]

    return {
        "ir_version": IR_VERSION,
        "source": {
            "excel_file": str(path),
            "excel_name": path.name,
            "workbook_version": workbook_version,
            "fingerprint": workbook_fingerprint(path, workbook_version),
            "catalog_source_dir": str(catalog_dir) if catalog_dir else None,
        },
        "targets": {
            "process_name": process_name,
            "project_name": project_name,
        },
        "process": {
            "work_item_types": work_item_types,
            "fields": fields,
            "html_fields": html_fields,
            "picklists": picklists,
            "backlogs": backlogs,
        },
        "project": {
            "teams": teams,
            "area_paths": areas,
            "iteration_paths": iterations,
            "project_placeholders_replaced": placeholders_replaced,
        },
        "catalog": catalog,
        "counts": {
            "work_item_types": len(enabled_wits),
            "disabled_work_item_types": len(disabled_wits),
            "fields": len(fields),
            "html_fields": len(html_fields),
            "hidden_field_placements": sum(
                len(wit.get("hidden_field_labels") or []) for wit in enabled_wits
            ),
            "picklists": len(picklists),
            "picklist_values": sum(len(values) for values in picklists.values()),
            "teams": len(teams),
            "area_paths": len(areas),
            "iteration_paths": len(iterations),
            "backlogs": len(backlogs),
            "catalog_files": len(catalog["files"]),
            "catalog_rows": catalog["total_rows"],
        },
    }
